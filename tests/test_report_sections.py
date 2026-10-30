import base64
import io
import json
import unittest
from openpyxl import load_workbook
from core import now
from reports import REPORT_GROUPS, SECTIONS
import test_miniapp as fixtures

class ReportSectionTests(unittest.TestCase):
    setUp=fixtures.MiniAppTests.setUp
    tearDown=fixtures.MiniAppTests.tearDown

    def report(self,section,uid=10,key=None,start=None,end=None):
        date=now()[:10]
        key=key or ('section-report-'+section+'-0001')
        self.service.write(uid,key,'report',{'start':start or date,'end':end or date,'section':section})
        payloads=[json.loads(row['payload']) for row in self.s.rows('SELECT payload FROM outbox WHERE chat=? ORDER BY id',(uid,))]
        doc=next(row for row in reversed(payloads) if 'document' in row)
        return load_workbook(io.BytesIO(base64.b64decode(doc['document']))),doc,payloads

    def test_only_selected_sheets_are_exported_for_every_group_in_both_languages(self):
        for lang in ('uz','ru'):
            with self.s.db:self.s.db.execute('UPDATE users SET lang=? WHERE id=10',(lang,))
            for section,(_,fields) in REPORT_GROUPS.items():
                if section=='all':continue
                with self.subTest(lang=lang,section=section):
                    wb,doc,payloads=self.report(section,key=f'report-{lang}-{section}-0001')
                    expected=[SECTIONS[name][lang=='ru'][:31] for name in fields]
                    self.assertEqual(wb.sheetnames,expected)
                    self.assertIn('-'+section+'-',doc['filename'])

    def test_sales_scope_dates_and_idempotency_are_preserved(self):
        self.service.write(20,'section-own-sale-001','sell',self.sale_data)
        self.service.write(21,'section-other-sale1','sell',self.sale_data)
        wb,doc,_=self.report('sales',uid=20)
        self.assertEqual(wb['Sotuvlar'].max_row,2)
        before=len(self.s.rows('SELECT * FROM outbox'))
        self.report('sales',uid=20)
        self.assertEqual(len(self.s.rows('SELECT * FROM outbox')),before)
        wb,_,_=self.report('sales',key='section-empty-period',start='2020-01-01',end='2020-01-01')
        self.assertEqual(wb['Sotuvlar'].max_row,1)

    def test_debts_are_current_balances_for_sales_inside_selected_period(self):
        self.service.write(20,'section-debt-sale01','sell',self.sale_data)
        self.service.write(21,'section-debt-sale02','sell',self.sale_data)
        with self.s.db:self.s.db.execute("UPDATE sales SET created='2020-01-01T10:00:00+05:00' WHERE id=2")
        wb,_,_=self.report('debts')
        sheet=wb[SECTIONS['debts'][0]]
        self.assertEqual(sheet.max_row,2)
        headers=[cell.value for cell in sheet[1]]
        self.assertEqual(sheet.cell(2,headers.index('Qarz')+1).value,4000000)
        wb,_,_=self.report('debts',uid=21)
        self.assertEqual(wb[SECTIONS['debts'][0]].max_row,1)

    def test_current_stock_is_separate_from_historical_period(self):
        wb,_,payloads=self.report('stock',start='2020-01-01',end='2020-01-01')
        self.assertGreater(wb[SECTIONS['stock'][0]].max_row,1)
        self.assertTrue(any('Joriy holat' in row.get('text','') for row in payloads))
        self.assertFalse(any('Sotilgan summa' in row.get('text','') for row in payloads))

    def test_unknown_section_or_invalid_period_creates_no_output(self):
        for data in ({'section':'unknown','start':now()[:10],'end':now()[:10]},
                     {'section':'sales','start':'2027-01-01','end':'2026-01-01'}):
            before=len(self.s.rows('SELECT * FROM outbox'))
            with self.assertRaises(ValueError):
                self.service.write(10,'bad-section-key-001','report',data)
            self.assertEqual(len(self.s.rows('SELECT * FROM outbox')),before)

    def test_all_sections_remain_available(self):
        wb,_,_=self.report('all')
        for name in ('sales','payments','refunds','stock','warehouses','tasks','quotes'):
            self.assertIn(SECTIONS[name][0][:31],wb.sheetnames)
