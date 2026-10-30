import base64
import io
import json
import os
import tempfile
import unittest
from datetime import datetime, timedelta
from unittest.mock import patch
from openpyxl import load_workbook
from core import Store, money, TZ, now, parse_item_lines
from bot import Bot
from i18n import tr, label
from reports import workbook
from payments import normalize_payment, payment_labels, payment_label

class TestCRM(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.s=Store(os.path.join(self.tmp.name,'test.sqlite3'))
        self.s.configure([10,11],[20,21,22,23])
        self.b=Bot(self.s);self.update=0
        with self.s.db:
            self.point=self.s.add('points',name='Ombor 1')
            self.product=self.s.product(10,'Linglong 315/80 R22.5 TBR','2500000')
            self.client=self.s.client(20,'Transport MChJ','+998901234567','Toshkent','')
            self.s.move(20,self.product,self.point,10,'in','start')
        self.items=[dict(product=self.product,qty=4,price='2500000',discount=0)]

    def tearDown(self):
        self.s.db.close();self.tmp.cleanup()

    def msg(self,text,uid=20,**extra):
        self.update+=1
        update={'update_id':self.update,'message':{'from':{'id':uid,'first_name':f'User {uid}'},'chat':{'id':uid,'type':'private'},'text':text,**extra}}
        self.b.process(update)
        return update

    def sale(self,uid=20,paid='1000000'):
        with self.s.db:
            return self.s.order(uid,self.client,self.items,self.point,paid,'naqd')

    def stock(self):
        return self.s.stock(20)[0]['qty']

    def outputs(self,uid=20):
        return [json.loads(r['payload']) for r in self.s.rows('SELECT payload FROM outbox WHERE chat=? ORDER BY id',(uid,))]

    def dialog(self,cmd,values,uid=20):
        self.msg('/'+cmd,uid)
        for v in values:self.msg(str(v),uid)
        self.msg(tr('confirm',self.s.user(uid)['lang']),uid)

    def test_sale_total_and_balance(self):
        ref=self.sale()
        row=self.s.one('sales',ref)
        self.assertEqual(row['total'],1000000000)
        self.assertEqual(row['paid'],100000000)
        self.assertEqual(self.stock(),6)
        self.assertEqual(len(self.s.rows('SELECT * FROM payments')),1)

    def test_multiline_sale_rolls_back_if_second_product_short(self):
        items=self.items+[dict(product=self.product,qty=9,price='2500000',discount=0)]
        with self.assertRaises(ValueError):
            with self.s.db:self.s.order(20,self.client,items,self.point,'0','naqd')
        self.assertEqual(self.stock(),10)
        self.assertEqual(self.s.rows('SELECT * FROM sales'),[])
        self.assertEqual(self.s.rows("SELECT * FROM movements WHERE kind='sale'"),[])

    def test_quote_does_not_reserve_stock(self):
        with self.s.db:
            ref=self.s.order(20,self.client,self.items,quote=True,terms='3 kun')
        self.assertEqual(self.stock(),10)
        self.assertEqual(self.s.one('quotes',ref)['total'],1000000000)

    def test_payment_and_overpayment(self):
        ref=self.sale()
        with self.s.db:self.s.payment(21,ref,'500000','terminal')
        self.assertEqual(self.s.one('sales',ref)['paid'],150000000)
        with self.assertRaises(ValueError):
            with self.s.db:self.s.payment(21,ref,'999999999','naqd')
        self.assertEqual(self.s.one('sales',ref)['paid'],150000000)

    def test_staff_cannot_create_product(self):
        with self.assertRaises(PermissionError):self.s.product(20,'X','10')
        self.msg('/newproduct')
        self.assertIn('Kirish yopiq',self.outputs()[-1]['text'])

    def test_employee_cannot_writeoff(self):
        with self.assertRaises(PermissionError):
            with self.s.db:self.s.move(20,self.product,self.point,1,'writeoff','x')

    def test_director_writeoff_and_return(self):
        with self.s.db:
            self.s.move(10,self.product,self.point,2,'writeoff','damaged')
            self.s.move(20,self.product,self.point,1,'return','returned')
        self.assertEqual(self.stock(),9)

    def test_unknown_and_group_messages_denied(self):
        self.msg('/stock',99)
        self.assertIn('99',self.outputs(99)[-1]['text'])
        before=len(self.s.rows('SELECT * FROM outbox'))
        self.update+=1
        self.b.process({'update_id':self.update,'message':{'from':{'id':20},'chat':{'id':-100,'type':'group'},'text':'/stock'}})
        self.assertEqual(len(self.s.rows('SELECT * FROM outbox')),before)

    def test_user_caps_and_revocation(self):
        with self.assertRaises(ValueError):self.s.configure([10,11,12,13],[])
        self.s.configure([10],[20,21,22,23,24])
        self.assertTrue(self.s.user(24)['active'])
        with self.assertRaises(ValueError):self.s.configure([10],[10])
        self.s.configure([10],[20])
        with self.assertRaises(PermissionError):self.s.user(21)

    def test_duplicate_update_does_not_duplicate_shift(self):
        update=self.msg('/startwork')
        before=len(self.s.rows('SELECT * FROM outbox'))
        self.b.process(update)
        self.assertEqual(len(self.s.rows('SELECT * FROM shifts')),1)
        self.assertEqual(len(self.s.rows('SELECT * FROM outbox')),before)

    def test_shift_start_end_reports_director_notification(self):
        self.msg('/startwork')
        self.dialog('endwork',['5 qo‘ng‘iroq, 1 sotuv'])
        row=self.s.rows('SELECT * FROM shifts')[0]
        self.assertIsNotNone(row['ended'])
        self.assertEqual(row['report'],'5 qo‘ng‘iroq, 1 sotuv')
        self.assertTrue(any('Ish tugadi' in p.get('text','') for p in self.outputs(10)))
        self.assertTrue(any('document' in p for p in self.outputs()))

    def test_duplicate_start_and_end_without_start(self):
        self.msg('/startwork');self.msg('/startwork')
        self.assertEqual(len(self.s.rows('SELECT * FROM shifts')),1)
        with self.assertRaises(ValueError):
            with self.s.db:self.s.end(21,'x')

    def test_manual_fine_requires_director(self):
        with self.assertRaises(PermissionError):self.s.fine(20,21,'50000','x')
        self.dialog('fine',[20,'50000','Kechikish'],10)
        row=self.s.rows('SELECT * FROM fines')[0]
        self.assertEqual(row['approved_by'],10)
        self.assertEqual(row['amount'],5000000)

    def test_sale_full_dialog(self):
        self.dialog('sell',[self.client,self.point,f'{self.product} | 4 | 2500000','1000000','naqd'])
        self.assertEqual(self.stock(),6)
        self.assertIsNone(self.b.session(20))

    def test_cancel_does_not_write(self):
        self.msg('/sell');self.msg(str(self.client));self.msg('/cancel')
        self.assertIsNone(self.b.session(20));self.assertEqual(self.stock(),10)
        self.assertEqual(self.s.rows('SELECT * FROM sales'),[])

    def test_dialog_survives_restart(self):
        self.msg('/income');self.msg(str(self.product))
        self.b=Bot(self.s)
        self.assertEqual(self.b.session(20)['data']['product'],self.product)
        self.msg(str(self.point));self.msg('3');self.msg('-');self.msg(tr('confirm'))
        self.assertEqual(self.stock(),13)

    def test_readonly_lookup_during_dialog(self):
        self.msg('/sell');self.msg('/clients')
        self.assertEqual(self.b.session(20)['step'],0)
        self.assertIn('Transport',self.outputs()[-1]['text'])

    def test_russian_menu_and_report(self):
        self.msg('/language');self.msg('Русский')
        self.assertEqual(self.s.user(20)['lang'],'ru')
        self.msg(label('stock','ru'))
        self.assertIn('Linglong',self.outputs()[-1]['text'])
        self.msg('/report');self.msg(now()[:10]);self.msg(now()[:10]);self.msg(tr('confirm','ru'))
        self.assertTrue(any('Отчет:' in p.get('text','') for p in self.outputs()))

    def test_assigned_tasks_permissions_and_due(self):
        due=(datetime.now(TZ)+timedelta(days=1)).strftime('%Y-%m-%d %H:%M')
        with self.assertRaises(PermissionError):self.s.task(20,self.client,21,due,'x')
        with self.s.db:ref=self.s.task(10,self.client,21,due,'Call client')
        with self.assertRaises(PermissionError):self.s.complete_task(20,ref)
        with self.s.db:self.s.complete_task(21,ref)
        self.assertEqual(self.s.one('tasks',ref)['status'],'done')

    def test_reminders_queued_once_and_retried(self):
        with self.s.db:
            ref=self.s.add('tasks',client=self.client,actor=20,due='2020-01-01T10:00:00+05:00',note='Call',created=now())
        self.b.schedule();self.b.schedule()
        messages=self.outputs()
        self.assertEqual(len(messages),1)
        self.assertIn('ESLATMA',messages[0]['text'])
        class Failed:
            def call(self,*a):raise OSError()
        self.b.drain(Failed())
        self.assertEqual(self.s.rows('SELECT sent FROM outbox')[0]['sent'],0)
        class OK:
            def call(self,*a):return {}
        self.b.drain(OK())
        self.assertEqual(self.s.rows('SELECT sent FROM outbox')[0]['sent'],1)

    def test_daily_report_employee_scope(self):
        self.sale(20);self.sale(21)
        today=now()[:10]
        self.assertEqual(len(self.s.report(20,today,today)['sales']),1)
        self.assertEqual(len(self.s.report(10,today,today)['sales']),2)
        self.assertEqual(len(self.s.report(20,today,today)['sale_items']),1)

    def test_delivery_separate_fee_and_status(self):
        sale=self.sale()
        self.dialog('delivery',[self.client,sale,'Ali +998901112233','Sergeli','150000','100000'])
        r=self.s.rows('SELECT * FROM deliveries')[0]
        self.assertEqual(r['fee'],15000000)
        self.assertEqual(r['paid'],10000000)
        self.dialog('delivery_status',[r['id'],'delivered'])
        self.assertEqual(self.s.one('deliveries',r['id'])['status'],'delivered')
        with self.assertRaises(PermissionError):self.s.delivery_status(21,r['id'],'cancelled')

    def test_delivery_client_must_match_sale(self):
        sale=self.sale()
        with self.s.db:other=self.s.client(20,'B','+998909999999','x','')
        with self.assertRaises(ValueError):self.s.delivery(20,other,sale,'x','x','100','0')

    def test_call_and_client_history(self):
        self.dialog('call',[self.client,'qiziqdi','КП kerak'])
        data=self.s.client_card(20,self.client)
        self.assertEqual(data['calls'][0]['result'],'qiziqdi')

    def test_product_photo_dialog(self):
        self.msg('/newproduct',10);self.msg('Triangle 12.00R20',10);self.msg('3000000',10)
        self.msg('',10,photo=[{'file_id':'photo-file'}]);self.msg(tr('confirm'),10)
        rows=self.s.rows('SELECT * FROM products WHERE name=?',('Triangle 12.00R20',))
        self.assertEqual(rows[0]['photo'],'photo-file')

    def test_amount_validation(self):
        for value in ['-1','NaN','Infinity','1.001','1e100','abc']:
            with self.subTest(value=value):
                with self.assertRaises(ValueError):money(value)
        self.assertEqual(money('2 500,50'),250050)

    def test_xlsx_currency_and_formula_escape(self):
        data=workbook({'sales':[{'name':'=HYPERLINK("bad")','total':125050,'paid':0}]})
        wb=load_workbook(io.BytesIO(data));ws=wb.active
        self.assertEqual(ws.cell(2,2).value,1250.5)
        self.assertEqual(ws.cell(2,1).data_type,'s')
        self.assertTrue(ws.cell(2,1).value.startswith("'="))

    def test_no_automatic_fines(self):
        self.msg('/startwork')
        self.assertEqual(self.s.rows('SELECT * FROM fines'),[])

    def test_failed_confirm_keeps_inventory_and_session(self):
        self.dialog('sell',[self.client,self.point,f'{self.product} | 99 | 2500000','0','naqd'])
        self.assertEqual(self.stock(),10)
        self.assertEqual(self.b.session(20)['cmd'],'sell')
        self.assertEqual(self.s.rows('SELECT * FROM sales'),[])

    def test_report_range_validation(self):
        with self.assertRaises(ValueError):self.s.report(20,'2026-12-01','2026-01-01')

    def test_revoked_user_queued_data_is_not_delivered(self):
        with self.s.db:self.b.queue(21,'Sensitive data')
        self.s.configure([10],[20])
        class Recorder:
            calls=[]
            def call(self,*a):self.calls.append(a)
        api=Recorder();self.b.drain(api)
        self.assertEqual(api.calls,[])

    def test_quote_dialog_pdf_and_status(self):
        self.dialog('quote',[self.client,f'{self.product} | 2 | 2500000','NAQD','Olib ketish'])
        quotes=self.s.rows('SELECT * FROM quotes')
        self.assertEqual(len(quotes),1)
        docs=[base64.b64decode(p['document']) for p in self.outputs() if 'document' in p]
        self.assertTrue(docs[0].startswith(b'%PDF'))
        self.assertEqual(self.stock(),10)
        self.dialog('quote_status',[quotes[0]['id'],'sent'])
        self.assertEqual(self.s.one('quotes',quotes[0]['id'])['status'],'sent')

    def test_huge_order_rejected_without_overflow(self):
        with self.assertRaises(ValueError):
            with self.s.db:
                self.s.order(20,self.client,[dict(product=self.product,qty=1000000,price='1000000000000',discount=0)],self.point)
        self.assertEqual(self.stock(),10)

    def test_director_receives_staff_day_excel(self):
        self.msg('/startwork');self.dialog('endwork',['Ish tugadi'])
        self.assertTrue(any('document' in p for p in self.outputs(10)))
        self.assertTrue(any('document' in p for p in self.outputs(11)))

    def test_export_contains_client_names_and_debt(self):
        self.sale()
        data=self.b.enrich(self.s.report(20,now()[:10],now()[:10]))
        self.assertEqual(data['sales'][0]['client_name'],'Transport MChJ')
        self.assertEqual(data['sales'][0]['debt'],900000000)

    def test_item_separator_formats(self):
        for line in ['1|4|2400000','1-4-2400000','1 4 2400000',
                     ' 1 | 4 | 2400000 ', '1 - 4 - 2400000',
                     '1\t4\t2400000','1–4–2400000']:
            with self.subTest(line=line):
                self.assertEqual(parse_item_lines(line),[dict(product=1,qty=4,price='2400000',discount=0)])

    def test_item_price_and_multiline_compatibility(self):
        rows=parse_item_lines('\n1|4|2 400 000,50\n\n1-2-2400000.50\n1 1 2400000,50\n')
        self.assertEqual([money(r['price']) for r in rows],[240000050]*3)
        self.assertEqual([r['qty'] for r in rows],[4,2,1])

    def test_invalid_items_cannot_turn_into_valid_sales(self):
        for line in ['-1-4-2400000','1--4-2400000','1-4--2400000',
                     '1 4 -2400000','1|4|-2400000','1 4',
                     '1 4 2400000 0','1|4|2400000|4','1-4-2400000-0',
                     '1 4 2 400 000','1 0 2400000','']:
            with self.subTest(line=line):
                with self.assertRaises(ValueError):parse_item_lines(line)

    def test_new_formats_in_sale_and_quote_both_languages(self):
        for lang in ['uz','ru']:
            for cmd in ['sell','quote']:
                for sep in ['|','-',' ']:
                    with self.subTest(lang=lang,cmd=cmd,sep=sep):
                        with self.s.db:self.s.db.execute('UPDATE users SET lang=? WHERE id=20',(lang,))
                        self.msg('/'+cmd);self.msg(str(self.client))
                        if cmd=='sell':self.msg(str(self.point))
                        self.msg(sep.join([str(self.product),'1','2400000']))
                        if cmd=='sell':self.msg('0');self.msg('naqd')
                        else:self.msg('NAQD');self.msg('Olib ketish')
                        # Confirm preview uses the same totals before every save.
                        self.assertTrue(any('2 400 000.00' in p.get('text','') for p in self.outputs()))
                        self.msg(tr('confirm',lang))
                        self.assertIsNone(self.b.session(20))
        self.assertEqual(len(self.s.rows('SELECT * FROM sales')),6)
        self.assertEqual(len(self.s.rows('SELECT * FROM quotes')),6)
        self.assertEqual(self.stock(),4)
        for table in ['sales','quotes']:
            self.assertTrue(all(r['total']==240000000 for r in self.s.rows('SELECT total FROM '+table)))

    def test_payment_choices_and_aliases(self):
        for raw,expected in [('NAQD','naqd'),('NAXT','naqd'),('PERCHESLENIYA','perechisleniya'),
                             ('PERECHISLENIYA','perechisleniya'),('o‘tkazma','perechisleniya'),
                             ('KARTA ORQALI TO‘LOV','karta'),('terminal','karta'),
                             ('ПЕРЕЧИСЛЕНИЕ','perechisleniya'),('ОПЛАТА КАРТОЙ','karta')]:
            with self.subTest(raw=raw):self.assertEqual(normalize_payment(raw),expected)
        with self.assertRaises(ValueError):normalize_payment('other')

    def test_payment_buttons_in_both_dialogs_and_languages(self):
        sale=self.sale()
        for lang in ['uz','ru']:
            for cmd in ['sell','pay']:
                with self.subTest(lang=lang,cmd=cmd):
                    with self.s.db:self.s.db.execute('UPDATE users SET lang=? WHERE id=20',(lang,))
                    self.msg('/'+cmd)
                    if cmd=='sell':
                        for value in [self.client,self.point,f'{self.product} 1 2400000','0']:self.msg(str(value))
                    else:
                        self.msg(str(sale));self.msg('1')
                    keyboard=self.outputs()[-1]['reply_markup']['keyboard']
                    self.assertEqual([r[0]['text'] for r in keyboard[:-1]],payment_labels(lang))
                    self.msg(payment_labels(lang)[1]);self.msg(tr('confirm',lang))
                    self.assertIsNone(self.b.session(20))
        self.assertEqual(self.s.rows('SELECT payment FROM sales ORDER BY id DESC')[0]['payment'],'perechisleniya')
        self.assertEqual(self.s.rows('SELECT method FROM payments ORDER BY id DESC')[0]['method'],'perechisleniya')

    def test_invalid_payment_cannot_save_sale_or_change_balance(self):
        with self.assertRaises(ValueError):
            with self.s.db:self.s.order(20,self.client,self.items,self.point,'0','unsupported')
        self.assertEqual(self.stock(),10)
        self.assertEqual(self.s.rows('SELECT * FROM sales'),[])
        sale=self.sale()
        with self.assertRaises(ValueError):
            with self.s.db:self.s.payment(20,sale,'10','unsupported')
        self.assertEqual(self.s.one('sales',sale)['paid'],100000000)

    def test_payment_names_in_excel_preserve_historical_values(self):
        data=workbook({'sales':[dict(payment='terminal',total=10000)],'payments':[dict(method='perechisleniya',amount=10000)]},'ru')
        wb=load_workbook(io.BytesIO(data))
        self.assertEqual(wb.worksheets[0].cell(2,1).value,'ОПЛАТА КАРТОЙ')
        self.assertEqual(wb.worksheets[1].cell(2,1).value,'ПЕРЕЧИСЛЕНИЕ')
        self.assertEqual(payment_label('old custom method'),'old custom method')

    def test_new_orders_cannot_apply_discounts(self):
        lines=[dict(product=self.product,qty=1,price='2500000',discount=5)]
        for quote in [False,True]:
            with self.assertRaisesRegex(ValueError,'discount_removed'):
                with self.s.db:self.s.order(20,self.client,lines,self.point,quote=quote,payment='naqd')
        self.assertEqual(self.stock(),10)
        self.assertEqual(self.s.rows('SELECT * FROM sales'),[])
        self.assertEqual(self.s.rows('SELECT * FROM quotes'),[])

    def test_old_pending_discount_is_cancelled_without_saving(self):
        with self.s.db:self.b.set_session(20,{'cmd':'sell','step':5,'data':{'client':self.client,'point':self.point,'items':[dict(product=self.product,qty=4,price='2500000',discount=5)],'paid':'0','payment':'naqd'}})
        self.msg(tr('confirm'))
        self.assertIsNone(self.b.session(20))
        self.assertEqual(self.stock(),10)
        self.assertIn('Chegirma funksiyasi olib tashlandi',self.outputs()[-1]['text'])

    def test_historical_quote_pdf_keeps_stored_total(self):
        from quotation import quotation_pdf
        with self.s.db:
            ref=self.s.order(20,self.client,self.items,quote=True,terms='Old quote')
            self.s.db.execute('UPDATE quotes SET total=950000000 WHERE id=?',(ref,))
            self.s.db.execute('UPDATE quote_items SET discount=5 WHERE quote=?',(ref,))
        pdf=quotation_pdf(self.s,20,ref)
        self.assertTrue(pdf.startswith(b'%PDF'))
        self.assertEqual(self.s.one('quotes',ref)['total'],950000000)
        self.assertEqual(self.s.rows('SELECT discount FROM quote_items WHERE quote=?',(ref,))[0]['discount'],5)

    def test_discount_is_absent_from_preview_and_excel(self):
        self.msg('/sell');self.msg(str(self.client));self.msg(str(self.point))
        self.msg(f'{self.product}-1-2500000');self.msg('0');self.msg('NAQD')
        preview=self.outputs()[-1]['text']
        self.assertNotIn('%',preview)
        self.assertNotIn('Chegirma',preview)
        data=workbook({'sale_items':[dict(name='Old row',qty=4,price=250000000,discount=5)]})
        wb=load_workbook(io.BytesIO(data))
        self.assertNotIn('Chegirma, %',[c.value for c in wb.active[1]])

    def test_assortment_does_not_hide_products_without_points(self):
        with self.s.db:
            self.s.db.execute('DELETE FROM movements')
            self.s.db.execute('DELETE FROM stock')
            self.s.db.execute('DELETE FROM points')
        rows=self.s.stock(20)
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['name'],'Linglong 315/80 R22.5 TBR')
        self.assertIsNone(rows[0]['point_id'])
        self.msg('/stock')
        self.assertIn('Linglong',self.outputs()[-1]['text'])
        self.assertIn('RESTORE-DATA.bat',self.outputs()[-1]['text'])
        self.assertNotIn('Dona: 0',self.outputs()[-1]['text'])
        self.msg('/points')
        self.assertIn('Savdo nuqtasi hali yo‘q',self.outputs()[-1]['text'])

    def test_stock_search_sends_matching_photo_once_across_points(self):
        with self.s.db:
            self.s.db.execute('UPDATE products SET photo=? WHERE id=?',('large-photo-id',self.product))
            self.s.add('points',name='Ombor 2')
            self.s.product(10,'Other 385/65','3000000','other-photo')
        self.msg('/stock 315/80')
        photos=[p for p in self.outputs() if 'photo' in p]
        self.assertEqual(len(photos),1)
        self.assertEqual(photos[0]['photo'],'large-photo-id')
        self.assertIn('Linglong',photos[0]['caption'])
        self.assertIn('2 500 000',photos[0]['caption'])
        self.assertLess(len(photos[0]['caption']),1024)

    def test_stock_without_photo_or_matches_still_works(self):
        self.msg('/stock 315/80')
        self.assertIn('Linglong',self.outputs()[-1]['text'])
        self.assertFalse(any('photo' in p for p in self.outputs()))
        self.msg('/stock no-such-size')
        self.assertEqual(self.outputs()[-1]['text'],tr('empty','uz'))

    def test_quote_requires_selected_payment_and_fulfillment(self):
        self.dialog('quote',[self.client,f'{self.product} 4 2500000','PERECHISLENIYA','Yetkazib berish'])
        q=self.s.rows('SELECT * FROM quotes')[0]
        self.assertEqual(q['payment'],'perechisleniya')
        self.assertEqual(q['fulfillment'],'delivery')
        self.assertEqual(q['terms'],'')
        self.assertEqual(q['total'],1000000000)
        self.assertEqual(self.stock(),10)

    def test_quote_options_updates_existing_quote_only(self):
        with self.s.db:ref=self.s.order(20,self.client,self.items,quote=True,terms='10')
        self.dialog('quote_options',[ref,'KARTA ORQALI TO‘LOV','Olib ketish'])
        q=self.s.one('quotes',ref)
        self.assertEqual(q['payment'],'karta');self.assertEqual(q['fulfillment'],'pickup')
        self.assertEqual(q['total'],1000000000)
        self.assertEqual(len(self.s.rows('SELECT * FROM quotes')),1)
        self.assertEqual(self.stock(),10)
        with self.assertRaises(PermissionError):self.s.quote_options(21,ref,'NAQD','pickup')

    def test_russian_fulfillment_buttons_and_invalid_input(self):
        with self.s.db:self.s.db.execute("UPDATE users SET lang='ru' WHERE id=20")
        self.msg('/quote');self.msg(str(self.client));self.msg(f'{self.product} 1 2500000');self.msg('НАЛИЧНЫЕ')
        keyboard=self.outputs()[-1]['reply_markup']['keyboard']
        self.assertEqual([r[0]['text'] for r in keyboard[:-1]],['Самовывоз','Доставка'])
        self.msg('10')
        self.assertEqual(self.b.session(20)['step'],3)
        self.assertEqual(self.s.rows('SELECT * FROM quotes'),[])
        self.msg('Доставка');self.msg(tr('confirm','ru'))
        self.assertEqual(self.s.rows('SELECT * FROM quotes')[0]['fulfillment'],'delivery')

    def test_old_quote_session_reprompts_without_using_old_terms_as_payment(self):
        with self.s.db:self.b.set_session(20,{'cmd':'quote','step':3,'data':{'client':self.client,'items':self.items,'terms':'10'}})
        self.msg(tr('confirm'))
        self.assertEqual(self.b.session(20)['step'],2)
        self.assertNotIn('terms',self.b.session(20)['data'])
        self.msg('NAQD');self.msg('Olib ketish');self.msg(tr('confirm'))
        self.assertEqual(self.s.rows('SELECT * FROM quotes')[0]['fulfillment'],'pickup')

    def test_quote_schema_upgrade_keeps_existing_history(self):
        import sqlite3
        path=os.path.join(self.tmp.name,'old-schema.sqlite3')
        conn=sqlite3.connect(path)
        conn.execute("CREATE TABLE quotes(id INTEGER PRIMARY KEY,client INTEGER NOT NULL,actor INTEGER NOT NULL,total INTEGER NOT NULL,terms TEXT NOT NULL,lang TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'draft',created TEXT NOT NULL)")
        conn.execute("INSERT INTO quotes VALUES(1,1,20,123400,'10','uz','draft','2026-10-16')")
        conn.commit();conn.close()
        migrated=Store(path)
        q=migrated.one('quotes',1)
        self.assertEqual(q['total'],123400);self.assertEqual(q['terms'],'10')
        self.assertEqual(q['payment'],'');self.assertEqual(q['fulfillment'],'')
        migrated.db.close()

    def test_excel_replaces_employee_ids_at_original_column(self):
        row=dict(id=1,client=self.client,point=self.point,actor=20,total=1400000000,paid=140000000,payment='naqd')
        row['employee_name']='Sherzod'
        for lang,header in [('uz','Xodim'),('ru','Сотрудник')]:
            data=workbook({'sales':[row]},lang)
            wb=load_workbook(io.BytesIO(data));ws=wb.active
            heads=[c.value for c in ws[1]]
            self.assertEqual(heads[3],header)
            self.assertEqual(heads.count(header),1)
            self.assertEqual(ws.cell(2,4).value,'Sherzod')
            self.assertFalse(any('Xodim ID' in h or 'ID сотрудника' in h for h in heads))
        self.assertEqual(row['actor'],20)

    def test_names_in_all_operational_report_sections(self):
        with self.s.db:
            self.s.db.execute("UPDATE users SET name='Sherzod' WHERE id=20")
            self.s.db.execute("UPDATE users SET name='Avazjon' WHERE id=10")
        sections={name:[dict(id=1,actor=20)] for name in ['sales','sale_items','payments','calls','deliveries','shifts','quotes','tasks']}
        # The raw report's business fields are kept so enrich computes debt normally.
        sections['sales'][0].update(total=100,paid=0)
        sections['deliveries'][0].update(fee=100,paid=0)
        sections['fines']=[dict(id=1,actor=20,approved_by=10)]
        sections['client']=[dict(id=self.client,owner=20)]
        data=workbook(self.b.enrich(sections))
        wb=load_workbook(io.BytesIO(data))
        for ws in wb.worksheets:
            heads=[c.value for c in ws[1]]
            if ws.title=='Mijoz':
                self.assertEqual(ws.cell(2,heads.index('Mas’ul xodim')+1).value,'Sherzod')
            else:
                self.assertEqual(ws.cell(2,heads.index('Xodim')+1).value,'Sherzod')
            if ws.title=='Jarimalar':self.assertEqual(ws.cell(2,heads.index('Tasdiqlagan direktor')+1).value,'Avazjon')
            self.assertFalse(any(h in {'Xodim ID','Mas’ul ID','Tasdiqlagan direktor ID'} for h in heads))
        self.assertEqual(sections['sales'][0]['actor'],20)

    def test_unregistered_name_is_not_exported_as_numeric_id(self):
        # An invited user who has not pressed Start still has their numeric placeholder name.
        data=workbook(self.b.enrich({'shifts':[dict(id=1,actor=21)]}))
        wb=load_workbook(io.BytesIO(data));ws=wb.active
        self.assertEqual(ws.cell(2,2).value,'Ism kiritilmagan')
        self.assertNotEqual(ws.cell(2,2).value,'21')

    def test_excel_uses_name_received_from_telegram(self):
        self.msg('/start')
        self.sale()
        data=self.b.enrich(self.s.report(10,now()[:10],now()[:10]))
        wb=load_workbook(io.BytesIO(workbook(data)))
        ws=wb['Sotuvlar'];heads=[c.value for c in ws[1]]
        self.assertEqual(ws.cell(2,heads.index('Xodim')+1).value,'User 20')

if __name__=='__main__':unittest.main()
