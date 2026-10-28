import hashlib
import hmac
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
from urllib.parse import urlencode
from urllib.request import Request,urlopen
from urllib.error import HTTPError
from core import Store,now
from bot import Bot
from miniapp import authenticate,AuthError,MiniService,MiniServer,initialize

TOKEN='unit-test-secret-only'


def signed(uid=20,age=0,**extra):
    fields={'auth_date':str(int(time.time())-age),'query_id':'fixture','user':json.dumps({'id':uid,'first_name':'Tester'},separators=(',',':')),**extra}
    check='\n'.join(f'{k}={v}' for k,v in sorted(fields.items()))
    secret=hmac.new(b'WebAppData',TOKEN.encode(),hashlib.sha256).digest()
    fields['hash']=hmac.new(secret,check.encode(),hashlib.sha256).hexdigest()
    return urlencode(fields)


class MiniAppTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.path=Path(self.temp.name)/'test.sqlite3'
        self.s=Store(str(self.path));self.s.configure([10],[20,21]);initialize(self.path)
        self.service=MiniService(self.s,TOKEN)
        with self.s.db:
            self.point=self.s.add('points',name='Main')
            self.client=self.s.client(20,'Customer','+998901234567','','')
            self.product=self.s.product(10,'Linglong 315/80 R22.5','2500000')
            self.s.move(20,self.product,self.point,10,'in','')
        self.sale_data={'client':self.client,'point':self.point,'items':[{'product':self.product,'qty':2,'price':'2500000'}],'paid':'1000000','payment':'naqd'}

    def tearDown(self):self.s.db.close();self.temp.cleanup()

    def test_excel_report_queues_real_workbook_once_with_employee_scope(self):
        import base64,io
        from openpyxl import load_workbook
        self.service.write(20,'report-sale-own-001','sell',self.sale_data)
        self.service.write(21,'report-sale-other-01','sell',self.sale_data)
        data={'start':now()[:10],'end':now()[:10]}
        self.service.write(20,'report-request-0001','report',data)
        self.service.write(20,'report-request-0001','report',data)
        docs=[json.loads(r['payload']) for r in self.s.rows('SELECT payload FROM outbox WHERE chat=20') if 'document' in json.loads(r['payload'])]
        self.assertEqual(len(docs),1)
        self.assertTrue(docs[0]['filename'].endswith('.xlsx'))
        wb=load_workbook(io.BytesIO(base64.b64decode(docs[0]['document'])))
        self.assertEqual(wb['Sotuvlar'].max_row,2)
        before=len(self.s.rows('SELECT * FROM outbox'))
        with self.assertRaises(ValueError):self.service.write(20,'report-bad-date-01','report',{'start':'bad','end':'bad'})
        self.assertEqual(len(self.s.rows('SELECT * FROM outbox')),before)

    def test_new_product_with_warehouse_opening_stock_is_atomic(self):
        data={'name':'Test tyre','brand':'New brand','size':'12 R20','price':'100','point':self.point,'qty':'3'}
        ref=self.service.write(10,'product-stock-0001','product',data)['id']
        self.assertEqual(self.s.db.execute('SELECT qty FROM stock WHERE product=?',(ref,)).fetchone()[0],3)
        self.service.write(10,'product-stock-0001','product',data)
        self.assertEqual(len(self.s.rows('SELECT * FROM movements WHERE product=?',(ref,))),1)
        before=len(self.s.rows('SELECT * FROM products'))
        with self.assertRaises(ValueError):self.service.write(10,'bad-product-qty-01','product',{**data,'qty':'1.5'})
        self.assertEqual(len(self.s.rows('SELECT * FROM products')),before)
        with self.assertRaises(ValueError):self.service.write(10,'missing-point-001','product',{**data,'point':''})
        self.assertEqual(len(self.s.rows('SELECT * FROM products')),before)

    def test_report_includes_movement_dates_names_and_current_warehouse_snapshot(self):
        with self.s.db:
            self.s.db.execute("UPDATE users SET name='Employee A' WHERE id=20")
            self.s.move(21,self.product,self.point,1,'in','Other employee')
            self.s.db.execute("UPDATE products SET price_pending=1 WHERE id=?",(self.product,))
        own=self.s.report(20,now()[:10],now()[:10])
        self.assertEqual(len(own['movements']),1)
        self.assertEqual(own['movements'][0]['warehouse_name'],'Main')
        self.assertTrue(own['movements'][0]['created'])
        director=self.s.report(10,now()[:10],now()[:10])
        self.assertEqual(len(director['movements']),2)
        self.assertEqual(len(director['product_events']),1)
        enriched=self.service.b.enrich(own)
        self.assertEqual(enriched['movements'][0]['employee_name'],'Employee A')
        stock=self.s.inventory_report(10)
        self.assertIsNone(stock['stock'][0]['price'])
        self.assertIsNone(stock['stock'][0]['stock_value'])
        self.assertEqual(stock['warehouses'][0]['qty'],11)
        self.assertEqual(stock['warehouses'][0]['unknown_price'],1)
        self.assertTrue(stock['warehouses'][0]['as_of'])

    def test_signed_auth_and_extra_signature_field(self):
        self.assertEqual(authenticate(signed(),TOKEN)['id'],20)
        self.assertEqual(authenticate(signed(signature='signed-extra'),TOKEN)['id'],20)

    def test_tampered_expired_future_and_duplicate_auth_are_rejected(self):
        cases=[signed().replace('Tester','Other'),signed(age=3601),signed(age=-61),signed()+'&auth_date=1','']
        for raw in cases:
            with self.assertRaises(AuthError):authenticate(raw,TOKEN)
        with self.assertRaises(AuthError):authenticate(signed(),'wrong-token')

    def test_idempotent_sale_and_payload_change_guard(self):
        first=self.service.write(20,'one-request-00000001','sell',self.sale_data)
        second=self.service.write(20,'one-request-00000001','sell',self.sale_data)
        self.assertEqual(first,second)
        self.assertEqual(self.s.stock(20)[0]['qty'],8)
        self.assertEqual(len(self.s.rows('SELECT * FROM payments')),1)
        with self.assertRaises(ValueError):self.service.write(20,'one-request-00000001','sell',{**self.sale_data,'paid':'2000000'})

    def test_invalid_stock_or_deadline_rolls_back_request_and_sale(self):
        for update in [{'items':[{'product':self.product,'qty':100,'price':'2500000'}]},{'debt_due':'bad-date'}]:
            with self.assertRaises(ValueError):self.service.write(20,'failed-request-0001','sell',{**self.sale_data,**update})
            self.assertEqual(self.s.rows('SELECT * FROM sales'),[])
            self.assertEqual(self.s.stock(20)[0]['qty'],10)
            self.assertEqual(self.s.rows('SELECT * FROM mini_requests'),[])

    def test_json_fractional_or_boolean_quantity_is_not_silently_rounded(self):
        for qty in [1.5,True,'1.5']:
            with self.assertRaises(ValueError):self.service.write(20,'invalid-qty-000001','sell',{**self.sale_data,'items':[{'product':self.product,'qty':qty,'price':'2500000'}]})
        self.assertEqual(self.s.rows('SELECT * FROM sales'),[])

    def test_api_roles_and_scoped_history(self):
        for cmd,data in [('product',{'name':'X','price':'1'}),('minimum',{'product':self.product,'point':self.point,'minimum':1}),('fine',{'employee':21,'amount':'1','reason':'X'})]:
            with self.assertRaises(PermissionError):self.service.write(20,'role-check-0000001',cmd,data)
        self.service.write(20,'sale-actor-0000001','sell',self.sale_data)
        self.assertEqual(len(self.service.history(20,'sales')),1)
        self.assertEqual(self.service.history(21,'sales'),[])
        self.assertEqual(len(self.service.history(10,'sales')),1)

    def test_shared_stock_and_money_between_miniapp_and_chat(self):
        self.service.write(20,'sale-shared-000001','sell',self.sale_data)
        other=Store(str(self.path));bot=Bot(other)
        try:
            self.assertEqual(other.stock(20)[0]['qty'],8)
            with other.db:other.payment(21,1,'4000000','karta')
            self.assertEqual(self.s.debts(10),[])
            self.assertEqual(self.service.dashboard(20)['received'],100000000)
            self.assertEqual(self.service.dashboard(10)['received'],500000000)
        finally:other.db.close()

    def test_quote_queues_pdf_without_stock_change(self):
        result=self.service.write(20,'quote-request-0001','quote',{'client':self.client,'items':self.sale_data['items'],'payment':'karta','fulfillment':'delivery'})
        self.assertEqual(result['id'],1);self.assertEqual(self.s.stock(20)[0]['qty'],10)
        self.assertIn('document',json.loads(self.s.rows('SELECT * FROM outbox')[0]['payload']))

    def test_large_telegram_ids_for_assignees(self):
        with self.s.db:
            self.s.db.execute("INSERT INTO users VALUES(1234567890,'employee','Person','uz',1)")
        self.service.write(10,'large-user-000001','task',{'client':self.client,'assignee':1234567890,'due':'2030-01-01T10:00','note':'Follow up'})
        self.assertEqual(self.s.one('tasks',1)['actor'],1234567890)

    def test_product_photo_permission_and_idempotent_write(self):
        with self.assertRaises(PermissionError):self.service.write(20,'photo-request-0001','product_photo',{'product':self.product,'image':'unused'})
        with patch.object(self.service,'upload_photo',return_value='verified-file-id') as upload:
            self.service.write(10,'photo-request-0001','product_photo',{'product':self.product,'image':'test'})
            self.service.write(10,'photo-request-0001','product_photo',{'product':self.product,'image':'test'})
            self.assertEqual(upload.call_count,1)
        self.assertEqual(self.s.one('products',self.product)['photo'],'verified-file-id')

    def test_launcher_button_only_with_configured_https_url(self):
        with patch.dict('os.environ',{'PUBLIC_APP_URL':'https://bot.example.uz'}):
            keyboard=self.service.b.keyboard(20)
            self.assertEqual(keyboard['keyboard'][0][0],{'text':'📱 Ilovani ochish'})
            for text in ['/app','📱 Ilovani ochish','📱 Открыть приложение']:
                self.service.b.route(20,{'text':text})
                payload=json.loads(self.s.rows('SELECT payload FROM outbox ORDER BY id DESC LIMIT 1')[0]['payload'])
                self.assertEqual(payload['reply_markup']['inline_keyboard'][0][0]['web_app']['url'],'https://bot.example.uz')
        with patch.dict('os.environ',{'PUBLIC_APP_URL':''}):
            self.assertFalse(any('web_app' in b for row in self.service.b.keyboard(20)['keyboard'] for b in row))

    def test_http_auth_origin_static_and_quote_ownership(self):
        self.service.write(20,'quote-http-000001','quote',{'client':self.client,'items':self.sale_data['items'],'payment':'naqd','fulfillment':'pickup'})
        server=MiniServer(('127.0.0.1',0),self.path,TOKEN,'https://bot.example.uz')
        thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        base=f'http://127.0.0.1:{server.server_port}'
        def request(path,uid=None,origin=None,data=None,key='http-request-0001'):
            headers={}
            if uid is not None:headers['X-Telegram-Init-Data']=signed(uid)
            if origin:headers['Origin']=origin
            if data:headers.update({'Content-Type':'application/json','Idempotency-Key':key})
            return urlopen(Request(base+path,headers=headers,data=json.dumps(data).encode() if data else None),timeout=5)
        try:
            with request('/') as response:self.assertIn(b'TOPSHINALAR',response.read())
            for uid,origin in [(None,None),(999,None),(20,'https://evil.example')]:
                with self.assertRaises(HTTPError) as error:request('/api/catalog',uid,origin)
                self.assertEqual(error.exception.code,403)
            with request('/api/catalog?q=315%2080',20) as response:
                self.assertEqual(json.load(response)['total'],1)
                self.assertEqual(response.headers['Cache-Control'],'no-store')
            for q in ['%20%20cUsToMeR%20%20','%2B998%20(90)%20123-45-67']:
                with request('/api/clients?q='+q,20) as response:self.assertEqual(json.load(response)['items'][0]['id'],self.client)
            with request('/api/clients?q=absent-client',20) as response:self.assertEqual(json.load(response)['total'],0)
            with request('/api/me',20) as response:self.assertEqual(json.load(response)['name'],'Tester')
            with self.assertRaises(HTTPError) as error:request('/api/quote/1.pdf',21)
            self.assertEqual(error.exception.code,403)
            with request('/api/quote/1.pdf',20) as response:self.assertTrue(response.read().startswith(b'%PDF'))
            body={'cmd':'sell','data':self.sale_data}
            with request('/api/action',20,data=body) as response:self.assertTrue(json.load(response)['ok'])
            with request('/api/action',20,data=body) as response:self.assertTrue(json.load(response)['ok'])
            self.assertEqual(self.s.stock(20)[0]['qty'],8)
            with self.s.db:self.s.db.execute('UPDATE users SET active=0 WHERE id=20')
            with self.assertRaises(HTTPError) as error:request('/api/me',20)
            self.assertEqual(error.exception.code,403)
        finally:server.shutdown();server.server_close();thread.join(timeout=5)
