import base64
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from bot import Bot
from core import Store, fmt
from i18n import tr
from usability import matches_product, size_parts
from backup import create_backup,verify_restore
from storage import migrate_database


class UsabilityTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name)/'topshinalar.sqlite3'
        self.s=Store(str(self.path));self.s.configure([10,11],[20,21])
        self.b=Bot(self.s);self.update=0
        with self.s.db:
            self.s.db.execute("UPDATE users SET name='Direktor' WHERE id=10")
            self.client=self.s.client(20,'Transport','+998901234567','Toshkent','')
            self.point=self.s.add('points',name='Asosiy ombor')
            self.product=self.s.product(10,'Linglong 315/80 R22.5 TBR','2500000','large-photo')
            self.s.move(20,self.product,self.point,20,'in','start')
        self.items=[dict(product=self.product,qty=2,price='2500000')]

    def tearDown(self):self.s.db.close();self.tmp.cleanup()

    def msg(self,text,uid=20):
        self.update+=1
        self.b.process({'update_id':self.update,'message':{'from':{'id':uid,'first_name':f'Person {uid}'},'chat':{'id':uid,'type':'private'},'text':str(text)}})

    def click(self,data,uid=20):
        self.update+=1
        update={'update_id':self.update,'callback_query':{'id':str(self.update),'from':{'id':uid,'first_name':f'Person {uid}'},'message':{'chat':{'id':uid,'type':'private'}},'data':data}}
        self.b.process(update);return update

    def outputs(self,uid=20):return [json.loads(r['payload']) for r in self.s.rows('SELECT payload FROM outbox WHERE chat=? ORDER BY id',(uid,))]

    def button(self,text,uid=20):
        for payload in reversed(self.outputs(uid)):
            for row in payload.get('reply_markup',{}).get('inline_keyboard',[]):
                for b in row:
                    if text in b['text']:return b['callback_data']
        self.fail(f'Button missing: {text}')

    def sale(self,paid='1000000',actor=20):
        with self.s.db:return self.s.order(actor,self.client,self.items,self.point,paid,'naqd')

    def qty(self):return self.s.db.execute('SELECT qty FROM stock WHERE product=? AND point=?',(self.product,self.point)).fetchone()[0]

    def test_plain_size_variants_show_same_card(self):
        for query in ['315/80','315 80','315/80R22.5','315/80 R22,5','linglong 315/80','LINGLONG']:
            self.msg(query)
            last=self.outputs()[-1]
            self.assertEqual(last['photo'],'large-photo',query)
            self.assertIn('Asosiy ombor: 20',last['caption'])
            self.assertIn('Linglong',last['caption'])
            self.assertIn('315/80',last['caption'])
        self.assertFalse(matches_product('315/80R24.5',{'name':'Linglong 315/80 R22.5'}))
        self.assertFalse(matches_product('385/65',{'name':'315/80 R22.5'}))

    def test_model_numbers_do_not_replace_tyre_size(self):
        for name, expected in [('Linglong KTD300 315/80 R22.5',('315','80','22.5')),('Triangle TR918 315/80 R22.5',('315','80','22.5')),('Bridgestone R249 385/65 R22.5',('385','65','22.5')),('Sailun S622 295/80 R22.5',('295','80','22.5'))]:
            self.assertEqual(size_parts(name),expected)
            self.assertTrue(matches_product(expected[0]+'/'+expected[1],{'name':name}))
        self.assertIsNone(size_parts('Triangle TR668 12.00 R20'))

    def test_catalog_pages_and_deduplicates_points(self):
        with self.s.db:
            self.s.add('points',name='Ikkinchi ombor')
            for i in range(6):self.s.product(10,f'Brand {i} 315/80 R22.5','1000','image-'+str(i))
        self.msg('315/80')
        self.assertEqual(sum('photo' in p for p in self.outputs()),4)
        self.click(self.button('Keyingi'))
        self.assertEqual(sum('photo' in p for p in self.outputs()),7)

    def test_card_actions_are_role_limited(self):
        self.msg('315/80')
        self.assertTrue(self.button('Sotish'))
        self.assertFalse(any('Tahrirlash' in str(p) for p in self.outputs()))
        self.msg('315/80',10)
        token=self.button('Tahrirlash',10)
        self.click(token,20)
        self.assertIsNone(self.b.session(20))
        self.assertIn(tr('stale'),self.outputs()[-1]['text'])
        self.msg('/editproduct',20)
        self.assertIn('Kirish yopiq',self.outputs()[-1]['text'])

    def test_quick_sale_uses_names_buttons_cart_and_confirmation(self):
        self.msg('315 80');self.click(self.button('Sotish'))
        self.click(self.button('Transport'))
        self.click(self.button('Asosiy ombor'))
        self.assertEqual(self.b.session(20)['item_stage'],'qty')
        self.msg('2');self.msg('2500000')
        self.click(self.button('Mahsulotlar tayyor'))
        self.msg('1000000');self.msg('NAQD')
        self.assertIn('Qolgan qarz: 4 000 000.00',self.outputs()[-1]['text'])
        self.assertEqual(self.qty(),20)
        self.msg(tr('confirm'))
        self.assertEqual(self.qty(),18)
        self.assertEqual(self.s.one('sales',1)['paid'],100000000)
        self.assertIsNone(self.b.session(20))
        self.click(self.button('To‘lov muddati'))
        self.msg('2027-01-10');self.msg(tr('confirm'))
        self.assertEqual(self.s.one('sales',1)['debt_due'],'2027-01-10')

    def test_quick_quote_and_cart_add_more_leave_stock_unchanged(self):
        self.msg('/stock');self.click(self.button('КП tayyorlash'))
        self.click(self.button('Transport'));self.msg('1');self.msg('2500000')
        self.click(self.button('Yana mahsulot'))
        self.msg('315 80');self.click(self.button('Linglong'))
        self.msg('2');self.msg('2500000');self.click(self.button('Mahsulotlar tayyor'))
        self.msg('KARTA ORQALI TO‘LOV');self.msg('Yetkazib berish');self.msg('1');self.msg(tr('confirm'))
        self.assertEqual(self.s.one('quotes',1)['total'],750000000)
        self.assertEqual(self.qty(),20)
        self.assertTrue(any('document' in p for p in self.outputs()))

    def test_picker_search_page_and_stale_click(self):
        with self.s.db:
            for i in range(9):self.s.client(20,f'Customer {i}','+998901234567','','')
        self.msg('/sell');old=self.button('Customer 0')
        self.click(self.button('Keyingi'))
        self.assertEqual(self.b.session(20)['step'],0)
        self.msg('Transport');token=self.button('Transport')
        self.click(token)
        self.assertEqual(self.b.session(20)['step'],1)
        self.click(old)
        self.assertEqual(self.b.session(20)['step'],1)
        self.assertIn(tr('stale'),self.outputs()[-1]['text'])

    def test_selector_callback_survives_restart_and_update_is_idempotent(self):
        self.msg('/sell');token=self.button('Transport');self.b=Bot(self.s)
        update=self.click(token)
        self.assertEqual(self.b.session(20)['step'],1)
        self.b.process(update)
        self.assertEqual(self.b.session(20)['step'],1)

    def test_back_preserves_earlier_fields_and_cancel_never_saves(self):
        self.msg('/sell');self.msg(self.client);self.msg(self.point)
        self.msg(f'{self.product} 2 2500000');self.msg('1000000');self.msg('naqd')
        self.msg(tr('back'))
        state=self.b.session(20)
        self.assertEqual(state['step'],4)
        self.assertEqual(state['data']['paid'],'1000000')
        self.assertNotIn('payment',state['data'])
        self.msg('/cancel')
        self.assertEqual(self.qty(),20)
        self.assertEqual(self.s.rows('SELECT * FROM sales'),[])

    def test_cart_back_edits_previous_item_instead_of_duplicating(self):
        self.msg('/stock');self.click(self.button('Sotish'))
        self.msg(self.client);self.msg(self.point);self.msg('2');self.msg('2500000')
        self.msg(tr('back'));self.msg('2400000')
        self.assertEqual(len(self.b.session(20)['cart']),1)
        self.assertEqual(self.b.session(20)['cart'][0]['price'],'2400000')

    def test_edit_product_keeps_history_and_stock(self):
        sale=self.sale('0');old=self.s.rows('SELECT * FROM sale_items WHERE sale=?',(sale,))[0]
        self.msg('/editproduct',10)
        for value in [self.product,'New name','New brand','315/80 R22.5','2700000','-']:
            self.msg(value,10)
        self.msg(tr('confirm'),10)
        row=self.s.one('products',self.product)
        self.assertEqual(row['price'],270000000)
        self.assertEqual(row['photo'],'large-photo')
        self.assertEqual(row['brand'],'New brand')
        self.assertEqual(self.qty(),18)
        self.assertEqual(self.s.rows('SELECT * FROM sale_items WHERE sale=?',(sale,))[0],old)

    def test_cancel_sale_restores_stock_and_records_refund_once(self):
        sale=self.sale();self.s.payment(21,sale,'500000','karta');self.s.db.commit()
        self.msg('/cancel_sale',10)
        for value in [sale,'Noto‘g‘ri sotuv',tr('refund_yes')]:self.msg(value,10)
        self.assertEqual(self.qty(),18)
        self.msg(tr('confirm'),10)
        self.assertEqual(self.qty(),20)
        self.assertEqual(self.s.one('sales',sale)['status'],'cancelled')
        self.assertEqual(sum(p['amount'] for p in self.s.rows('SELECT * FROM payments')),150000000)
        self.assertEqual(self.s.rows('SELECT * FROM refunds')[0]['amount'],150000000)
        self.assertEqual(self.s.debts(10),[])
        with self.assertRaises(ValueError):
            with self.s.db:self.s.cancel_sale(10,sale,'Again',True)
        self.assertEqual(self.qty(),20)
        with self.assertRaises(ValueError):
            with self.s.db:self.s.payment(20,sale,'1','naqd')

    def test_cancel_requires_director_reason_and_refund_confirmation(self):
        sale=self.sale()
        for uid,reason,confirmed,error in [(20,'reason',True,PermissionError),(10,'',True,ValueError),(10,'reason',False,ValueError)]:
            with self.assertRaises(error):
                with self.s.db:self.s.cancel_sale(uid,sale,reason,confirmed)
        self.assertEqual(self.qty(),18)

    def test_cancellation_refuses_changed_payment_after_preview(self):
        sale=self.sale()
        self.msg('/cancel_sale',10)
        for value in [sale,'Mistake',tr('refund_yes')]:self.msg(value,10)
        with self.s.db:self.s.payment(21,sale,'1000','naqd')
        self.msg(tr('confirm'),10)
        self.assertEqual(self.s.one('sales',sale)['status'],'active')
        self.assertIn('to‘lov o‘zgargan',self.outputs(10)[-1]['text'])
        self.msg(tr('back'),10);self.msg(tr('refund_yes'),10);self.msg(tr('confirm'),10)
        self.assertEqual(self.s.rows('SELECT * FROM refunds')[0]['amount'],100100000)

    def test_cancellation_rolls_back_if_stock_restore_fails(self):
        sale=self.sale()
        self.s.db.execute("CREATE TRIGGER fail_restore BEFORE UPDATE ON stock WHEN NEW.qty>OLD.qty BEGIN SELECT RAISE(ABORT,'test failure'); END")
        with self.assertRaises(sqlite3.IntegrityError):
            with self.s.db:self.s.cancel_sale(10,sale,'Mistake',True)
        self.assertEqual(self.s.one('sales',sale)['status'],'active')
        self.assertEqual(self.s.rows('SELECT * FROM refunds'),[])
        self.assertEqual(self.qty(),18)

    def test_refunds_in_reports_and_cancelled_debt_is_zero(self):
        sale=self.sale()
        with self.s.db:self.s.cancel_sale(10,sale,'Mistake',True)
        self.msg('/report',10);self.msg('2020-01-01',10);self.msg('2030-01-01',10);self.msg(tr('confirm'),10)
        documents=[p for p in self.outputs(10) if 'document' in p]
        self.assertEqual(len(documents),1)
        wb=load_workbook(io.BytesIO(base64.b64decode(documents[0]['document'])))
        self.assertIn('Pul qaytarishlar',wb.sheetnames)
        sheet=wb['Sotuvlar'];headers=[c.value for c in sheet[1]]
        self.assertEqual(sheet.cell(2,headers.index('Qarz')+1).value,0)
        self.assertIn('Bekor qilgan direktor',headers)
        messages='\n'.join(p.get('text','') for p in self.outputs(10))
        self.assertIn('Sotilgan summa: 0.00',messages)
        self.assertIn('Sof pul tushumi: 0.00',messages)

    def test_minimum_alert_once_rearms_after_restock_and_can_disable(self):
        with self.s.db:self.s.set_minimum(10,self.product,self.point,20)
        self.b.schedule();self.b.schedule()
        alerts=lambda:sum('KAM QOLDIQ' in p.get('text','') for p in self.outputs(10))
        self.assertEqual(alerts(),1)
        with self.s.db:self.s.move(20,self.product,self.point,2,'in','restock')
        self.b.schedule()
        with self.s.db:self.s.move(10,self.product,self.point,2,'writeoff','used')
        self.b.schedule();self.assertEqual(alerts(),2)
        with self.s.db:self.s.set_minimum(10,self.product,self.point,0)
        self.b.schedule();self.assertEqual(alerts(),2)

    def test_threshold_for_empty_point_and_employee_permissions(self):
        with self.s.db:
            empty=self.s.add('points',name='Empty point')
            self.s.set_minimum(10,self.product,empty,3)
        self.b.schedule()
        self.assertTrue(any('Empty point: 0 / min 3' in p.get('text','') for p in self.outputs(10)))
        with self.assertRaises(PermissionError):self.s.set_minimum(20,self.product,empty,1)

    def test_debt_deadline_scope_and_overdue(self):
        sale=self.sale();other=self.sale(actor=21)
        with self.s.db:self.s.set_debt_due(20,sale,'2020-01-01')
        with self.assertRaises(PermissionError):
            with self.s.db:self.s.set_debt_due(21,sale,'2020-01-01')
        self.assertEqual([r['id'] for r in self.s.debts(20,True)],[sale])
        self.assertEqual(len(self.s.debts(10)),2)
        self.msg('/overdue');self.assertIn('2020-01-01',self.outputs()[-1]['text'])
        self.msg('/debts',21);self.assertIn(f'#{other}',self.outputs(21)[-1]['text'])

    def test_deadline_is_stored_in_sortable_date_format(self):
        sale=self.sale()
        with self.s.db:self.s.set_debt_due(20,sale,'2027-1-2')
        self.assertEqual(self.s.one('sales',sale)['debt_due'],'2027-01-02')

    def test_due_reminders_once_daily_and_no_revoked_actor(self):
        sale=self.sale()
        with self.s.db:
            self.s.set_debt_due(20,sale,'2020-01-01')
            self.s.db.execute('UPDATE users SET active=0 WHERE id=20')
        with patch('usability.now',return_value='2026-10-24T10:00:00+05:00'):
            self.b.schedule();self.b=Bot(self.s);self.b.schedule()
        self.assertEqual(sum('TO‘LOV ESLATMASI' in p.get('text','') for p in self.outputs(10)),1)
        self.assertFalse(any('TO‘LOV ESLATMASI' in p.get('text','') for p in self.outputs(20)))
        with patch('usability.now',return_value='2026-10-25T10:00:00+05:00'):self.b.schedule()
        self.assertEqual(sum('TO‘LOV ESLATMASI' in p.get('text','') for p in self.outputs(10)),2)

    def test_fully_paid_debt_never_reminds(self):
        sale=self.sale()
        with self.s.db:
            self.s.set_debt_due(20,sale,'2020-01-01')
            self.s.payment(20,sale,'4000000','naqd')
        with patch('usability.now',return_value='2026-10-24T10:00:00+05:00'):self.b.schedule()
        self.assertFalse(any('TO‘LOV ESLATMASI' in p.get('text','') for p in self.outputs(10)))

    def test_daily_summary_is_once_per_director_and_day(self):
        with patch('usability.now',return_value='2026-10-24T01:00:00+05:00'):
            self.b.schedule();self.b=Bot(self.s);self.b.schedule()
        messages=[p['text'] for p in self.outputs(10) if 'KUNLIK XULOSA' in p.get('text','')]
        self.assertEqual(len(messages),1)
        self.assertIn('2026-10-23',messages[0])
        self.msg('/summary',20);self.assertIn('Kirish yopiq',self.outputs()[-1]['text'])

    def test_backup_restore_comparison_and_resumed_business_operations(self):
        sale=self.sale();self.msg('/sell');self.msg(self.client)
        with self.s.db:
            self.s.set_debt_due(20,sale,'2027-01-01')
            self.s.set_minimum(10,self.product,self.point,5)
        backup=create_backup(self.path,Path(self.tmp.name)/'backups')
        counts=verify_restore(backup)
        self.assertEqual(counts['sales'],1);self.assertEqual(counts['sessions'],1)
        target=Path(self.tmp.name)/'restored'/'topshinalar.sqlite3'
        migrate_database(backup,target)
        restored=Store(str(target))
        try:
            self.assertEqual(restored.one('sales',sale)['debt_due'],'2027-01-01')
            self.assertEqual(Bot(restored).session(20)['data']['client'],self.client)
            with restored.db:restored.payment(20,sale,'4000000','naqd')
            self.assertEqual(restored.debts(10),[])
            self.assertEqual(self.s.one('sales',sale)['paid'],100000000)
        finally:restored.db.close()

    def test_restore_foreign_key_corruption_is_detected(self):
        sale=self.sale();backup=create_backup(self.path,Path(self.tmp.name)/'backups')
        with sqlite3.connect(backup) as c:c.execute('UPDATE sales SET client=999 WHERE id=?',(sale,))
        with self.assertRaises(ValueError):verify_restore(backup)

    def test_new_schema_survives_reopen_and_existing_data_is_unchanged(self):
        sale=self.sale();before=self.s.one('sales',sale)
        other=Store(str(self.path))
        try:self.assertEqual(other.one('sales',sale),before)
        finally:other.db.close()

    def test_russian_catalog_cart_back_and_debt_views(self):
        with self.s.db:self.s.db.execute("UPDATE users SET lang='ru' WHERE id=20")
        self.msg('315 80');self.assertIn('Размер: 315/80 R22.5',self.outputs()[-1]['caption'])
        self.click(self.button('Продать'));self.click(self.button('Transport'));self.click(self.button('Asosiy'))
        self.msg('1');self.msg(tr('back','ru'))
        self.assertEqual(self.b.session(20)['item_stage'],'qty')
        self.msg('/cancel');self.sale();self.msg('/debts')
        self.assertIn('Срок не назначен',self.outputs()[-1]['text'])

    def test_photo_card_is_sent_with_inline_markup(self):
        self.msg('315/80')
        payload=self.outputs()[-1]
        class FakeAPI:
            def __init__(self):self.requests=[]
            def call(self,method,data):self.requests.append((method,data))
        api=FakeAPI();self.b.drain(api)
        method,data=api.requests[-1]
        self.assertEqual(method,'sendPhoto')
        self.assertEqual(data['photo'],'large-photo')
        self.assertIn('inline_keyboard',data['reply_markup'])

    def test_back_on_first_step_returns_to_menu(self):
        self.msg('/sell');self.msg(tr('back'))
        self.assertIsNone(self.b.session(20))
        self.assertIn('keyboard',self.outputs()[-1]['reply_markup'])

    def test_paged_task_button_opens_confirm_and_marks_only_own_task(self):
        with self.s.db:
            for i in range(7):self.s.task(20,self.client,20,'2030-01-01 10:00',f'Task {i}')
        self.msg('/tasks');self.assertTrue(self.button('Keyingi'))
        self.click(self.button('Bajarildi'));self.msg(tr('confirm'))
        self.assertEqual(len(self.s.rows("SELECT * FROM tasks WHERE status='done'")),1)
        self.msg('/tasks',21);self.assertEqual(self.outputs(21)[-1]['text'],tr('empty'))

    def test_old_v1_schema_migration_preserves_sale_stock_and_permissions(self):
        from core import SCHEMA
        path=Path(self.tmp.name)/'old-v1.sqlite3'
        with sqlite3.connect(path) as c:
            c.executescript(SCHEMA)
            c.execute("INSERT INTO users(id,role,name) VALUES(10,'director','Director')")
            c.execute("INSERT INTO clients VALUES(1,'Old client','123','','',10,'2020-01-01')")
            c.execute("INSERT INTO points VALUES(1,'Old point')")
            c.execute("INSERT INTO products(id,name,price) VALUES(1,'Old tire',10000)")
            c.execute('INSERT INTO stock VALUES(1,1,8)')
            c.execute("INSERT INTO sales(id,client,point,actor,total,paid,payment,created) VALUES(1,1,1,10,20000,5000,'naqd','2020-01-01')")
            c.execute("INSERT INTO sale_items VALUES(1,1,1,'Old tire',2,10000,0)")
        migrated=Store(str(path))
        try:
            self.assertEqual(migrated.one('sales',1)['paid'],5000)
            self.assertEqual(migrated.one('sales',1)['status'],'active')
            self.assertEqual(migrated.one('sales',1)['debt_due'],'')
            self.assertEqual(migrated.stock(10)[0]['qty'],8)
            self.assertEqual(migrated.user(10)['role'],'director')
            self.assertEqual(migrated.one('products',1)['brand'],'')
        finally:migrated.db.close()

    def test_refund_on_later_day_is_accounted_on_refund_date(self):
        with patch('core.now',return_value='2026-10-23T10:00:00+05:00'):sale=self.sale()
        with patch('core.now',return_value='2026-10-24T10:00:00+05:00'):
            with self.s.db:self.s.cancel_sale(10,sale,'Mistake',True)
        first=self.s.report(10,'2026-10-23','2026-10-23')
        second=self.s.report(10,'2026-10-24','2026-10-24')
        self.assertEqual(sum(p['amount'] for p in first['payments']),100000000)
        self.assertEqual(first['refunds'],[])
        self.assertEqual(second['payments'],[])
        self.assertEqual(sum(p['amount'] for p in second['refunds']),100000000)
