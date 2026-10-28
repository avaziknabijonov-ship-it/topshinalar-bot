"""Director-only, reviewed XLSX stock receipts. No writes during preview."""
import base64
import hashlib
import io
import json
import re
import zipfile
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook
from core import money, now

class ExcelError(ValueError):
    pass

def fail(uz,ru):
    raise ExcelError('Excel: '+uz+' / '+ru)

def normalized(value):
    return re.sub(r'[^\w]','',str(value or '').casefold())

ALIASES={
 'id':['id','product id','mahsulot id','код товара'],
 'sku':['sku','артикул','catalog sku'],
 'name':['name','nom','nomi','mahsulot','mahsulot nomi','tovar','tovar nomi','название','название товара','наименование','наименование товара','модель','model'],
 'brand':['brand','brend','бренд','марка'],
 'size':['size','razmer','размер'],
 'price':['price','narx','narxi','цена','цена uzs','narx uzs','цена продажи'],
 'qty':['kelgan dona','kelgan soni','kirim','приход','поступило','qty','quantity','qoldiq','ostatok','dona','soni','остаток','количество','колво','остаток шт','остаток на складе','остаток на точке','остатки'],
}
HEADERS={normalized(alias):key for key,aliases in ALIASES.items() for alias in aliases}

def number(value,qty=False):
    if isinstance(value,bool) or value is None:raise ValueError()
    raw=re.sub(r'[\s\u00a0\u202f]','',str(value)).replace(',','.')
    if isinstance(value,str) and re.search(r'[,.]\d{3}$',raw):raise ValueError()
    val=Decimal(raw)
    if not val.is_finite() or val<0 or val>(1000000 if qty else 1000000000000):raise ValueError()
    if qty:
        if val!=val.to_integral_value():raise ValueError()
        return int(val)
    if val.as_tuple().exponent < -2:raise ValueError()
    return money(str(val))

class InventoryExcel:
    def __init__(self,store):self.s=store

    def preview(self,uid,encoded,point,sheet=''):
        self.s.director(uid);self.s.one('points',point)
        try:
            if not isinstance(encoded,str) or len(encoded)>5600000:raise ValueError()
            raw=base64.b64decode(encoded,validate=True)
            if not 0<len(raw)<=4*1024*1024:raise ValueError()
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist())>1000 or sum(i.file_size for i in archive.infolist())>24*1024*1024:raise ValueError()
            wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        except Exception:fail('Fayl .xlsx bo‘lsin, hajmi 4 MB gacha.','Нужен файл .xlsx размером до 4 МБ.')
        try:
            if sheet and sheet not in wb.sheetnames:fail('Varaq topilmadi.','Лист не найден.')
            ws=wb[sheet] if sheet else wb.active
            if ws.max_column and ws.max_column>100:fail('Ustunlar soni 100 dan oshmasin.','Не более 100 столбцов.')
            cols=None;records=[]
            for n,row in enumerate(ws.iter_rows(values_only=True),1):
                if n>2020:fail('Bir importda 2000 qatorgacha.','Не более 2000 строк за импорт.')
                if not any(v is not None and str(v).strip() for v in row):continue
                if cols is None:
                    found={}
                    for i,v in enumerate(row):
                        key=HEADERS.get(normalized(v))
                        if key:
                            if key in found:fail('Bir xil ustun sarlavhasi takrorlangan.','Повторяется заголовок столбца.')
                            found[key]=i
                    if {'price','qty'}<=found.keys() and {'id','sku','name'} & found.keys():cols=found;continue
                    if n>=20:break
                    continue
                if any(isinstance(v,str) and v.startswith('=') for v in row):fail(f'{n}-qator: formula o‘rniga qiymat yozing.',f'Строка {n}: замените формулы значениями.')
                data={key:row[i] if i<len(row) else None for key,i in cols.items()}
                try:price=number(data.get('price'));qty=number(data.get('qty'),True)
                except (ValueError,InvalidOperation,TypeError):fail(f'{n}-qator: narx yoki qoldiq noto‘g‘ri.',f'Строка {n}: некорректная цена или остаток.')
                def text(key):
                    val=str(data.get(key) or '').strip()
                    if len(val)>200:fail(f'{n}-qator: matn juda uzun.',f'Строка {n}: слишком длинный текст.')
                    return val
                name,brand,size,sku=text('name'),text('brand'),text('size'),text('sku')
                if data.get('id') not in (None,''):
                    try:ref=number(data['id'],True);old=self.s.one('products',ref);matches=[old]
                    except (ValueError,InvalidOperation,TypeError):fail(f'{n}-qator: mahsulot ID topilmadi.',f'Строка {n}: ID товара не найден.')
                elif sku:
                    matches=self.s.rows('SELECT * FROM products WHERE catalog_sku=?',(sku,))
                else:
                    if not name:fail(f'{n}-qator: mahsulot nomi kerak.',f'Строка {n}: требуется название товара.')
                    matches=[p for p in self.s.rows('SELECT * FROM products') if normalized(p['name'])==normalized(name) and (not brand or normalized(p['brand'])==normalized(brand)) and (not size or normalized(p['size'])==normalized(size))]
                if len(matches)>1:fail(f'{n}-qator: bir nechta mahsulot topildi, ID yoki SKU kiriting.',f'Строка {n}: несколько совпадений; укажите ID или SKU.')
                old=matches[0] if matches else None
                if old and not old['active']:fail(f'{n}-qator: mahsulot faol emas.',f'Строка {n}: товар неактивен.')
                if not old and not name:fail(f'{n}-qator: yangi mahsulot nomi kerak.',f'Строка {n}: требуется название нового товара.')
                stock=self.s.db.execute('SELECT qty FROM stock WHERE product=? AND point=?',(old['id'],point)).fetchone() if old else None
                records.append({'row':n,'id':old['id'] if old else None,'name':old['name'] if old else name,'brand':old['brand'] if old else brand,'size':old['size'] if old else size,'sku':sku if not old else old['catalog_sku'],'old_qty':stock[0] if stock else None,'old_price':old['price'] if old else None,'old_price_pending':old['price_pending'] if old else None,'qty':qty,'new_qty':(stock[0] if stock else 0)+qty,'price':price})
                if len(records)>2000:fail('2000 qatordan ko‘p.','Более 2000 строк.')
            if cols is None:fail('Sarlavha: Nomi (yoki ID/SKU), Narx, Kelgan dona.','Заголовки: Название (или ID/SKU), Цена, Количество.')
            if not records:fail('Mahsulot qatorlari yo‘q.','Нет строк товаров.')
            seen=set()
            for r in records:
                key=('id',r['id']) if r['id'] else ('sku',r['sku']) if r['sku'] else ('new',normalized(r['name']),normalized(r['brand']),normalized(r['size']))
                if key in seen:fail(f"{r['row']}-qator: mahsulot takrorlangan.",f"Строка {r['row']}: повтор товара.")
                seen.add(key)
            fingerprint=hashlib.sha256(json.dumps(['excel_receipt_v2',point,ws.title,records],sort_keys=True).encode()).hexdigest()
            return {'sheet':ws.title,'sheets':wb.sheetnames,'count':len(records),'new_count':sum(r['id'] is None for r in records),'fingerprint':fingerprint,'rows':records}
        finally:wb.close()

    def apply(self,uid,encoded,point,sheet,fingerprint):
        preview=self.preview(uid,encoded,point,sheet)
        if not isinstance(fingerprint,str) or preview['fingerprint']!=fingerprint:fail('Qoldiq yoki narx o‘zgardi. Faylni qayta tekshiring.','Остаток или цена изменились. Повторите проверку файла.')
        for row in preview['rows']:
            ref=row['id']
            if ref is None:
                ref=self.s.add('products',name=row['name'],brand=row['brand'],size=row['size'],price=row['price'],catalog_sku=row['sku'] or None,price_pending=0)
                self.s.log(uid,'product',ref)
            else:
                self.s.db.execute('UPDATE products SET price=?,price_pending=0 WHERE id=?',(row['price'],ref))
                self.s.log(uid,'product_edit',ref)
            self.s.db.execute('INSERT INTO stock(product,point,qty) VALUES(?,?,?) ON CONFLICT(product,point) DO UPDATE SET qty=excluded.qty',(ref,point,row['new_qty']))
            self.s.add('movements',product=ref,point=point,qty=row['qty'],kind='excel_income',actor=uid,note=f"Excel {preview['sheet']} row {row['row']}: {row['old_qty'] or 0} + {row['qty']} = {row['new_qty']}",created=now())
            self.s.rearm_stock(ref,point)
        self.s.log(uid,'excel_import',point)
        return {'count':preview['count'],'new_count':preview['new_count']}
