import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from payments import payment_label, fulfillment_label

COLUMNS={
'currency':('Valyuta','Валюта'),'exchange_rate':('USD kursi, so‘m','Курс USD, сум'),'source_total_usd':('Asl jami, USD','Исходная сумма, USD'),
'warehouse_name':('Ombor','Склад'),'stock_value':('Narxi ma’lum qoldiq qiymati','Стоимость остатков с ценой'),'unknown_stock':('Qoldig‘i kiritilmagan mahsulotlar','Товары без остатка'),'unknown_price':('Narxi kiritilmagan qoldiqlar','Остатки без цены'),'as_of':('Qoldiq olingan vaqt','Время снимка остатков'),'stock_status':('Qoldiq holati','Статус остатка'),'kind':('Harakat turi','Тип движения'),'action':('Amal','Действие'),
'brand':('Brend','Бренд'),'size':('Razmer','Размер'),'debt_due':('To‘lov muddati','Срок оплаты'),
'cancel_reason':('Bekor qilish sababi','Причина отмены'),'cancelled_by':('Bekor qilgan direktor ID','ID отменившего директора'),'cancelled_by_name':('Bekor qilgan direktor','Отменивший директор'),'cancelled_at':('Bekor qilingan vaqt','Время отмены'),
'id':('ID','ID'),'name':('Nomi / ism','Название / имя'),'phone':('Telefon','Телефон'),'address':('Manzil','Адрес'),
'note':('Izoh','Примечание'),'owner':('Mas’ul ID','ID ответственного'),'client':('Mijoz ID','ID клиента'),
'client_name':('Mijoz','Клиент'),'actor':('Xodim ID','ID сотрудника'),'employee_name':('Xodim','Сотрудник'),
'owner_name':('Mas’ul xodim','Ответственный сотрудник'),'approved_by_name':('Tasdiqlagan direktor','Утвердивший директор'),
'point':('Nuqta','Точка'),'point_id':('Nuqta ID','ID точки'),'product':('Mahsulot ID','ID товара'),
'price':('Narx','Цена'),'qty':('Dona','Количество'),'discount':('Chegirma, %','Скидка, %'),
'sale':('Sotuv ID','ID продажи'),'quote':('КП ID','ID КП'),'total':('Jami','Итого'),
'paid':('To‘langan','Оплачено'),'debt':('Qarz','Долг'),'payment':('To‘lov turi','Тип оплаты'),
'amount':('Summa','Сумма'),'method':('To‘lov turi','Тип оплаты'),'created':('Sana va vaqt','Дата и время'),
'due':('Eslatma vaqti','Время напоминания'),'status':('Holat','Статус'),'notified':('Eslatma yuborildi','Напоминание отправлено'),
'result':('Natija','Результат'),'started':('Ish boshlandi','Начало работы'),'ended':('Ish tugadi','Конец работы'),
'report':('Kun hisoboti','Отчет за день'),'reason':('Sabab','Причина'),'approved_by':('Tasdiqlagan direktor ID','ID утвердившего директора'),
'driver':('Haydovchi','Водитель'),'fee':('Yetkazish narxi','Стоимость доставки'),
'fulfillment':('Shartlar','Условия'),'lang':('Til','Язык'),'terms':('Avvalgi izoh','Прежнее примечание'),'role':('Lavozim','Роль'),
}
SECTIONS={
'movements':('Ombor harakatlari','Движения склада'),'product_events':('Mahsulot qo‘shish tarixi','История товаров'),'warehouses':('Omborlar hozirgi qoldiq','Склады текущие остатки'),
'refunds':('Pul qaytarishlar','Возвраты денег'),
'sales':('Sotuvlar','Продажи'),'sale_items':('Sotilgan mahsulotlar','Проданные товары'),
'payments':('To‘lovlar','Платежи'),'calls':('Qo‘ng‘iroqlar','Звонки'),'deliveries':('Labo yetkazish','Доставка Labo'),
'shifts':('Ish kunlari','Рабочие дни'),'fines':('Jarimalar','Штрафы'),'quotes':('Tijorat takliflari','Коммерческие предложения'),
'tasks':('Vazifalar','Задачи'),'client':('Mijoz','Клиент'),'stock':('Qoldiqlar','Остатки'),
}
VALUES={'director':('Direktor','Директор'),'employee':('Xodim','Сотрудник'),'open':('Ochiq','Открыта'),'done':('Bajarilgan','Выполнена'),'draft':('Qoralama','Черновик'),'sent':('Yuborilgan','Отправлено'),'accepted':('Qabul qilingan','Принято'),'rejected':('Rad etilgan','Отклонено'),'new':('Yangi','Новая'),'on_way':('Yo‘lda','В пути'),'delivered':('Yetkazildi','Доставлено'),'cancelled':('Bekor qilingan','Отменено')}

def localized(key,val,lang):
    if key in {'kind','action','stock_status'}:
        labels={'in':('Kirim','Приход'),'opening':('Boshlang‘ich qoldiq','Начальный остаток'),'return':('Qaytarish','Возврат'),'writeoff':('Hisobdan chiqarish','Списание'),'sale':('Sotuv','Продажа'),'product':('Mahsulot qo‘shildi','Товар добавлен'),'product_edit':('Mahsulot tahrirlandi','Товар изменён'),'unknown':('Kiritilmagan','Не указан'),'known':('Kiritilgan','Указан')}
        if val in labels:return labels[val][lang=='ru']
    if key=='status' and val=='active':return 'Faol' if lang=='uz' else 'Активна'
    if key=='method' and val=='refund':return 'Pul qaytarish' if lang=='uz' else 'Возврат денег'
    if key in {'payment','method'}:return payment_label(val,lang)
    if key=='fulfillment':return fulfillment_label(val,lang)
    if key in {'status','role'} and val in VALUES:return VALUES[val][lang=='ru']
    return val

MONEY={'price','total','paid','amount','fee','debt','stock_value','source_total_usd'}
PEOPLE={'actor':'employee_name','owner':'owner_name','approved_by':'approved_by_name','cancelled_by':'cancelled_by_name'}

def export_fields(rows):
    """Replace person identifiers in-place and omit duplicate name columns."""
    fields=[]
    for key in rows[0] if rows else ['id']:
        if key=='discount':continue
        key=PEOPLE.get(key,key)
        if key not in fields:fields.append(key)
    return fields

def workbook(sections,lang='uz'):
    wb=Workbook();wb.remove(wb.active)
    for index,(name,rows) in enumerate(sections.items()):
        ws=wb.create_sheet(SECTIONS.get(name,(name,name))[lang=='ru'][:31])
        defaults={'movements':['created','name','brand','size','warehouse_name','qty','kind','actor','note'],'product_events':['created','name','brand','size','action','actor'],'warehouses':['warehouse_name','qty','stock_value','unknown_stock','unknown_price','as_of']}
        fields=export_fields(rows if rows else [dict.fromkeys(defaults[name])]) if name in defaults else export_fields(rows)
        ws.append([COLUMNS.get(k,(k,k))[lang=='ru'] for k in fields])
        for row in rows:
            values=[]
            for key in fields:
                val=localized(key,row.get(key),lang)
                if key in PEOPLE.values() and not val:
                    val='Имя не указано' if lang=='ru' else 'Ism kiritilmagan'
                if key in MONEY and val is not None:val=val/100
                if isinstance(val,str) and val.lstrip().startswith(('=','+','-','@')):val="'"+val
                values.append(val)
            ws.append(values)
        for c in ws[1]:
            c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='17394D');c.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[1].height=34;ws.freeze_panes='A2'
        for idx,key in enumerate(fields,1):
            ws.column_dimensions[get_column_letter(idx)].width=40 if key in {'name','note','address','report','terms',*PEOPLE.values()} else 23
            for cells in ws.iter_rows(min_row=2,min_col=idx,max_col=idx):
                c=cells[0];c.alignment=Alignment(wrap_text=True,vertical='top')
                if key in MONEY:c.number_format='#,##0.00'
        if rows:
            tab=Table(displayName=f'Data{index}',ref=f'A1:{get_column_letter(len(fields))}{len(rows)+1}')
            tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            ws.add_table(tab)
        ws.sheet_view.showGridLines=False
    if not sections:wb.create_sheet('TOPSHINALAR')
    out=io.BytesIO();wb.save(out);return out.getvalue()
